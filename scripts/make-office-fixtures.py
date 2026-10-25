"""Writes the Word, Excel and PowerPoint files the browser tests convert.
Run once (python3 scripts/make-office-fixtures.py); the files are committed.
Needs: pip install python-docx openpyxl python-pptx, and LibreOffice (to fill in formula results)."""
import io, os, subprocess, tempfile, datetime
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_BREAK
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt
from pptx.dml.color import RGBColor

OUT = 'e2e/fixtures/office'
os.makedirs(OUT, exist_ok=True)

def png(width, height, rgb):
    """A plain-colour PNG, written by hand so no imaging library is needed."""
    import zlib, struct
    raw = b''.join(b'\x00' + bytes(rgb) * width for _ in range(height))
    def chunk(t, d): return struct.pack('>I', len(d)) + t + d + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', width, height, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')

# Word: headings, styled runs, lists, a table, a picture and a page break.
doc = Document()
doc.add_heading('Quarterly Report', level=1)
p = doc.add_paragraph('Sales grew by ')
p.add_run('24 percent').bold = True
p.add_run(' this quarter, led by the ')
p.add_run('western region').italic = True
p.add_run('.')
doc.add_heading('Highlights', level=2)
for item in ['New store in Pune', 'Online orders doubled', 'Costs down 8%']:
    doc.add_paragraph(item, style='List Bullet')
for item in ['Hire two analysts', 'Open the Nagpur office']:
    doc.add_paragraph(item, style='List Number')
table = doc.add_table(rows=3, cols=3)
table.style = 'Table Grid'
for r, row in enumerate([['Region', 'Q1', 'Q2'], ['West', '120', '149'], ['North', '98', '103']]):
    for c, value in enumerate(row):
        table.cell(r, c).text = value
doc.add_picture(io.BytesIO(png(120, 60, (255, 170, 0))), width=Inches(2))
doc.add_paragraph('Long paragraph. ' * 60)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading('Appendix', level=1)
doc.add_paragraph('Prepared by Happy PDF tests. Price: ₹1,250.')
doc.save(f'{OUT}/report.docx')

# Excel: two sheets with text, numbers, a date, a percentage and a formula (results filled in by LibreOffice).
wb = Workbook()
ws = wb.active
ws.title = 'Sales'
ws.append(['Region', 'Units', 'Price', 'Total', 'Share', 'Date'])
for c in ws[1]: c.font = Font(bold=True)
rows = [('West', 120, 49.5), ('North', 98, 52.25), ('South', 143, 47.0), ('East', 77, 55.75)]
for i, (region, units, price) in enumerate(rows, start=2):
    ws.append([region, units, price, f'=B{i}*C{i}', f'=B{i}/SUM($B$2:$B$5)', datetime.date(2026, 4, i)])
    ws[f'C{i}'].number_format = '#,##0.00'
    ws[f'D{i}'].number_format = '#,##0.00'
    ws[f'E{i}'].number_format = '0.0%'
    ws[f'F{i}'].number_format = 'dd-mm-yyyy'
ws.column_dimensions['A'].width = 14
notes = wb.create_sheet('Notes')
notes['A1'] = 'Figures are in rupees.'
with tempfile.TemporaryDirectory() as tmp:
    raw = os.path.join(tmp, 'sales.xlsx')
    wb.save(raw)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', OUT, raw], check=True, capture_output=True)

# PowerPoint: a title slide, a bullet slide, and a slide with a picture and a coloured box.
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Happy PDF Launch'
s.placeholders[1].text = 'Every document tool. Free. Private.'
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = 'Why people switch'
body = s.placeholders[1].text_frame
body.text = 'Nothing is uploaded'
for line in ['Works offline', 'No sign-up, no limits']:
    body.add_paragraph().text = line
s = prs.slides.add_slide(prs.slide_layouts[6])
box = s.shapes.add_shape(1, PInches(0.5), PInches(0.5), PInches(4), PInches(1))
box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xFF, 0xC2, 0x4D)
box.text_frame.text = 'Picture slide'
s.shapes.add_picture(io.BytesIO(png(200, 100, (30, 120, 200))), PInches(5), PInches(2), PInches(4), PInches(2))
tb = s.shapes.add_textbox(PInches(0.5), PInches(5), PInches(8), PInches(1))
tb.text_frame.text = 'Made with python-pptx'
tb.text_frame.paragraphs[0].runs[0].font.size = PPt(28)
prs.save(f'{OUT}/launch.pptx')
print('wrote', sorted(os.listdir(OUT)))
